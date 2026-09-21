

from aitestapp.scripts.vectorIndexing import VectorIndexing
from aitestapp.scripts.sentencesearch import SimilarSentenceSearch
from aitestapp.scripts.config_functest import loadConfig
from aitestapp.scripts.model import ZhiPu4
from aitestapp.scripts.check import fix_parentheses
import re, time, ast, json, logging, os

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import html as html_module

logger = logging.getLogger(__name__)

class runner():
    def __init__(self):
        self.config = loadConfig()
        self.configs = self.config.load_config()

        self.SUMMARY_LENGTH = self.configs['SUMMARY_LENGTH']

    def running_1(self, prd_name):

        # 执行拆分需求
        # 执行拆分需求转为向量
        # 将向量写入向量数据库
        # 将向量数据与文本数据通过MySQL关联
        vector_indexing = VectorIndexing(prd_name)
        vector_indexing.process_texts()

    def running_2(self, prd_name):
        # 获取所有文档的热词
        # 从数据库中查询所有的文本
        # 使用 settings 或配置文件中的路径，避免硬编码
        faiss_index_path = os.path.join(self.configs.get('FAISS_BASE_DIR', '.'), self.configs['FAISS_DB_PATH'])
        similar_sentence_search = SimilarSentenceSearch(faiss_index_path)
        zhipu_4 = ZhiPu4()

        information_document = similar_sentence_search.get_information_document()
        return information_document
        information_document_list = zhipu_4.split_list_by_text_length(information_document, self.SUMMARY_LENGTH)
        reponse = zhipu_4.concurrent_zhipuai_request_text(information_document_list)
        logger.info(f"AI响应类型: {type(reponse)}")

        need_text_list = []
        for data in reponse:
            test = re.sub(r'```', '', data)
            test = re.sub(r'python', '', test)
            need_text_list.append(test)

        # 相似度阈值
        similarity_threshold = self.configs['SIMILARITY_THRESHOLD']
        vdb = similar_sentence_search.vectordb

        # 创建Word文档 - 使用配置路径或 settings
        doc_dir = self.configs.get('DOCUMENT_DIR', 'document')
        doc_filename = os.path.join(doc_dir, prd_name)
        doc = Document()
        doc.add_heading('Requirements', 0)  # 添加标题

        # 设置样式
        style = doc.styles['Normal']
        font = style.font
        font.size = Pt(12)
        font.name = 'Arial'

        elements = []
        for data in need_text_list:
            logger.debug(f"处理数据: {data}, 类型: {type(data)}")
            # 使用安全的 JSON 解析替代 eval()
            try:
                parsed_data = json.loads(data)
                if isinstance(parsed_data, list):
                    text_list = parsed_data
                elif isinstance(parsed_data, dict):
                    text_list = [parsed_data]
                else:
                    logger.warning(f"无法解析的数据格式: {type(parsed_data)}")
                    continue
            except (json.JSONDecodeError, TypeError):
                # 如果不是JSON格式，尝试作为纯文本处理
                logger.warning(f"数据非JSON格式，尝试作为文本处理: {str(data)[:100]}")
                try:
                    # 尝试使用 ast.literal_eval（比 eval 安全，但只能处理基本类型）
                    text_list = ast.literal_eval(data)
                    if not isinstance(text_list, (list, tuple)):
                        text_list = [text_list]
                except (ValueError, SyntaxError):
                    # 都失败则跳过该条数据
                    logger.error(f"数据解析完全失败，跳过: {str(data)[:100]}")
                    continue

            for text_item in text_list:
                # 根据实际数据结构调整，这里假设是字符串或字典
                if isinstance(text_item, str):
                    text = text_item
                elif isinstance(text_item, dict):
                    # 如果是字典，提取关键字段
                    text = text_item.get('text', text_item.get('content', str(text_item)))
                else:
                    text = str(text_item)

                logger.info(f"处理文本: {text}")
                # 执行相似度查询
                similar_texts = similar_sentence_search.search(text, similarity_threshold)
                logger.info(f"相似度搜索结果: {similar_texts}")

                if similar_texts is not None:
                    # 将数据存入数据库
                    vdb.create_needs_text(text, str(similar_texts))

                    # 将HTML实体转换回原始字符
                    text = html_module.unescape(text)
                    similar_texts = html_module.unescape(str(similar_texts))

                    # 添加文本到Word文档
                    doc.add_paragraph(text, style=style)
                    doc.add_paragraph(str(similar_texts), style=style)

                    # 添加一些空白
                    doc.add_paragraph('', style=style)

        # 保存Word文档
        os.makedirs(os.path.dirname(doc_filename) or '.', exist_ok=True)
        doc.save(doc_filename)
        logger.info(f"文档已保存: {doc_filename}")
from bs4 import BeautifulSoup
import requests
import json
from datetime import datetime
from docx import Document
from docx.shared import Inches
# 发送请求获取Swagger数据
# url = 'https://api.apiopen.top/swagger/doc.json'  # Swagger JSON数据的URL
# response = requests.get(url, verify=False)

# # html_content = response.content
# # html_dict = json.loads(html_content)
# # print(type(html_dict['paths']))
# # print(html_dict['paths'])
# html_content = response.text

# # 使用BeautifulSoup解析
# soup = BeautifulSoup(html_content, 'html.parser')
# # p_tags = soup.find_all('div')
# # # print(soup)
# # print(p_tags)
# span = soup.find_all('div', class_='opblock-summary opblock-summary-post')
# print(span)
# description = soup.find_all(string='responses')
# print(description)


# swagger_data = response.json()
# print(swagger_data)
# # 解析Swagger数据
# paths = swagger_data.get('paths', {})
# print(paths)

 
# # 生成带时间戳的文件名
# current_time = datetime.now().strftime('%Y%m%d%H%M%S')
# file_name = f"VisionCube_API_{current_time}.txt"
 

url = 'https://api.apiopen.top/swagger/doc.json'  # Swagger JSON数据的URL
# 启用 SSL 证书验证，防止中间人攻击（生产环境必须启用）
response = requests.get(url, verify=True, timeout=10)
swagger_data = response.json()
 
# 解析Swagger数据
paths = swagger_data.get('paths', {})
 
# 生成带时间戳的文件名
current_time = datetime.now().strftime('%Y%m%d%H%M%S')
file_name = f"开放接口{current_time}.docx"
# 创建一个新的 DOCX 文件
doc = Document()

# 接口数据
for path, methods in paths.items():
    for method, info in methods.items():
        # doc.add_paragraph("{)
        doc.add_paragraph('example: {} {}\n'.format(method,path))
        doc.add_paragraph('url: {}\n'.format(path))
        doc.add_paragraph('Method: {}\n'.format(method))
        doc.add_paragraph('notes: {}\n'.format(info.get('summary', '')))
        doc.add_paragraph('title: {}\n'.format(info.get('description', '')))
        doc.add_paragraph('Params: {}\n'.format(info.get('parameters', [])))
        doc.add_paragraph('Responses: [{}]\n'.format(info.get('responses', {})))
        status_text = ', '.join(["['{}', '{}']".format(code, info['responses'][code]['description']) for code in info['responses']])
        doc.add_paragraph('status: {}\n'.format(status_text))
        # doc.add_paragraph("},\n")
# 保存 DOCX 文件
doc.save(file_name)
print(f"API接口信息已保存到{file_name}")
